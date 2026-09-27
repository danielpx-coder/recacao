"""Testes de importação e exportação (TXT, DOCX, PDF, WordPress HTML, ZIP e acentuação brasileira)."""

import zipfile
from pathlib import Path
import pytest
from docx import Document
import pypdf

from danaurium.io.importers import import_txt, import_docx, import_pdf, import_file, ImportErrorWithDiagnostics
from danaurium.io.exporters import (
    sanitize_for_wordpress_html, export_txt, export_docx, export_markdown,
    export_wordpress_html, export_project_zip
)


def test_import_txt_utf8_and_brazilian_characters(tmp_path):
    txt_file = tmp_path / "materia_acentos.txt"
    content = "Atenção: Ações de educação pública em São Paulo e Brasília têm elevação de 15% no orçamento."
    with open(txt_file, "w", encoding="utf-8") as f:
        f.write(content)

    loaded = import_txt(txt_file)
    assert loaded == content
    assert "Atenção" in loaded
    assert "São Paulo" in loaded
    assert "orçamento" in loaded


def test_import_docx(tmp_path):
    docx_file = tmp_path / "documento_teste.docx"
    doc = Document()
    doc.add_paragraph("Primeiro parágrafo da reportagem oficial.")
    doc.add_paragraph("Segundo parágrafo com detalhes adicionais.")
    doc.save(str(docx_file))

    loaded = import_docx(docx_file)
    assert "Primeiro parágrafo da reportagem oficial." in loaded
    assert "Segundo parágrafo com detalhes adicionais." in loaded


def test_pdf_ocr_warning_when_empty_text_layer(tmp_path):
    # Criar PDF sem camada de texto (em branco / simulando escaneado)
    pdf_file = tmp_path / "escaneado_sem_ocr.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with open(pdf_file, "wb") as f:
        writer.write(f)

    with pytest.raises(ImportErrorWithDiagnostics) as exc_info:
        import_pdf(pdf_file)

    msg = str(exc_info.value)
    assert "OCR" in msg
    assert "camada de texto" in msg


def test_wordpress_html_sanitization():
    raw_draft = (
        "O Ministério da Educação anunciou novidades.\n\n"
        "> A inclusão digital é essencial para o país.\n\n"
        "- Primeiro benefício prático\n"
        "- Segundo benefício comprovado\n\n"
        "Mais detalhes em breve."
    )

    clean_html = sanitize_for_wordpress_html(raw_draft, title="Título da Matéria")

    # Verificações de segurança e semântica WordPress
    assert "<h2 class=\"wp-block-heading\">Título da Matéria</h2>" in clean_html
    assert "<p>O Ministério da Educação anunciou novidades.</p>" in clean_html
    assert "<blockquote class=\"wp-block-quote\"><p>A inclusão digital é essencial para o país.</p></blockquote>" in clean_html
    assert "<ul class=\"wp-block-list\">" in clean_html
    assert "<li>Primeiro benefício prático</li>" in clean_html
    assert "<li>Segundo benefício comprovado</li>" in clean_html

    # Garantir ausência de tags perigosas ou estrutura de página inteira
    assert "<script" not in clean_html
    assert "<html" not in clean_html
    assert "<body" not in clean_html
    assert "style=" not in clean_html


def test_export_project_zip_bundle(tmp_path):
    zip_path = tmp_path / "pacote_final.zip"
    products = {
        "reportagem": "Texto da reportagem completa com acentuação: ação, eleição.",
        "release": "Texto do release institucional.",
        "html_wordpress": "Texto para inserção no portal WordPress.",
    }

    out_zip = export_project_zip(zip_path, "Projeto Educação 2026", products)
    assert out_zip.exists()

    with zipfile.ZipFile(str(out_zip), "r") as zf:
        file_list = zf.namelist()
        assert "manifesto.txt" in file_list
        assert any("reportagem.txt" in name for name in file_list)
        assert any("reportagem.md" in name for name in file_list)
        assert any("html_wordpress.html" in name for name in file_list)

        # Verificar conteúdo com acentos dentro do ZIP
        rep_txt = zf.read([n for n in file_list if "reportagem.txt" in n][0]).decode("utf-8")
        assert "ação, eleição" in rep_txt
